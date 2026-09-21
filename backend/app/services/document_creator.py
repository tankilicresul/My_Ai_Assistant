import os
import json
import uuid
from typing import Dict, Any, List, Optional
from app.core.config import settings
from app.services.llm_gateway import llm_gateway

class DocumentCreatorService:
    """
    Service responsible for generating and building physical files on disk:
    - PDF (.pdf) using ReportLab
    - Word (.docx) using python-docx
    - Excel (.xlsx) using openpyxl & pandas
    - CSV (.csv) using pandas/csv
    - JSON (.json)
    - AI-assisted full document drafting & compilation
    """

    def create_pdf(self, title: str, sections: List[Dict[str, Any]], filename: Optional[str] = None) -> str:
        """
        Generate a formatted, professional PDF report.
        """
        fname = filename or f"generated_{uuid.uuid4().hex[:8]}.pdf"
        if not fname.endswith(".pdf"):
            fname += ".pdf"
        output_path = os.path.join(settings.DOCUMENT_STORAGE_DIR, fname)

        try:
            from reportlab.lib.pagesizes import letter, A4
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib import colors

            doc = SimpleDocTemplate(
                output_path,
                pagesize=A4,
                rightMargin=40,
                leftMargin=40,
                topMargin=40,
                bottomMargin=40
            )

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'DocTitle',
                parent=styles['Heading1'],
                fontSize=20,
                leading=24,
                textColor=colors.HexColor("#1e1b4b"),
                spaceAfter=15
            )
            heading_style = ParagraphStyle(
                'DocHeading',
                parent=styles['Heading2'],
                fontSize=14,
                leading=18,
                textColor=colors.HexColor("#4338ca"),
                spaceBefore=12,
                spaceAfter=8
            )
            body_style = ParagraphStyle(
                'DocBody',
                parent=styles['Normal'],
                fontSize=10,
                leading=14,
                textColor=colors.HexColor("#1e293b"),
                spaceAfter=8
            )

            story = []
            story.append(Paragraph(title, title_style))
            story.append(Spacer(1, 10))

            for sec in sections:
                sec_title = sec.get("title")
                sec_text = sec.get("content") or sec.get("text")
                sec_table = sec.get("table")

                if sec_title:
                    story.append(Paragraph(sec_title, heading_style))

                if sec_text:
                    for paragraph in sec_text.split("\n"):
                        if paragraph.strip():
                            story.append(Paragraph(paragraph.strip(), body_style))
                    story.append(Spacer(1, 6))

                if sec_table and isinstance(sec_table, list) and len(sec_table) > 0:
                    t = Table(sec_table)
                    t.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#4f46e5")),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, 0), 10),
                        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
                        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor("#f8fafc")),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ]))
                    story.append(t)
                    story.append(Spacer(1, 12))

            doc.build(story)
        except Exception as e:
            # Fallback text-based PDF representation if reportlab encounters layout error
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(f"%PDF-1.4\n1 0 obj <<\n/Title ({title})\n>>\n")
                f.write(f"\n{title}\n\n")
                for sec in sections:
                    f.write(f"## {sec.get('title', '')}\n{sec.get('content', '')}\n\n")

        return output_path

    def create_docx(self, title: str, sections: List[Dict[str, Any]], filename: Optional[str] = None) -> str:
        """
        Generate a Microsoft Word document (.docx).
        """
        fname = filename or f"generated_{uuid.uuid4().hex[:8]}.docx"
        if not fname.endswith(".docx"):
            fname += ".docx"
        output_path = os.path.join(settings.DOCUMENT_STORAGE_DIR, fname)

        try:
            import docx
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = docx.Document()

            # Document Title
            title_p = doc.add_heading(title, level=0)
            title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for sec in sections:
                sec_title = sec.get("title")
                sec_text = sec.get("content") or sec.get("text")
                sec_table = sec.get("table")

                if sec_title:
                    doc.add_heading(sec_title, level=1)

                if sec_text:
                    for para in sec_text.split("\n"):
                        if para.strip():
                            doc.add_paragraph(para.strip())

                if sec_table and isinstance(sec_table, list) and len(sec_table) > 0:
                    rows_count = len(sec_table)
                    cols_count = len(sec_table[0]) if rows_count > 0 else 0
                    if cols_count > 0:
                        tbl = doc.add_table(rows=rows_count, cols=cols_count)
                        tbl.style = 'Light Shading Accent 1'
                        for r_idx, row_data in enumerate(sec_table):
                            for c_idx, cell_value in enumerate(row_data):
                                tbl.cell(r_idx, c_idx).text = str(cell_value)
                        doc.add_paragraph()

            doc.save(output_path)
        except Exception as e:
            print(f"[DocumentCreator] Docx generation error: {e}")
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(f"{title}\n\n")
                for sec in sections:
                    f.write(f"# {sec.get('title', '')}\n{sec.get('content', '')}\n\n")

        return output_path

    def create_excel(self, title: str, columns: List[str], rows: List[List[Any]], filename: Optional[str] = None) -> str:
        """
        Generate an Excel spreadsheet (.xlsx) with styled header and auto column sizing.
        """
        fname = filename or f"generated_{uuid.uuid4().hex[:8]}.xlsx"
        if not fname.endswith(".xlsx"):
            fname += ".xlsx"
        output_path = os.path.join(settings.DOCUMENT_STORAGE_DIR, fname)

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = title[:30] if title else "Data"

            # Header Style
            header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            # Write Header
            for col_idx, col_name in enumerate(columns, 1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            # Write Data Rows
            for row_idx, row_data in enumerate(rows, 2):
                for col_idx, cell_value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
                    cell.border = thin_border
                    cell.font = Font(name="Calibri", size=10)

            # Auto fit column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    except:
                        pass
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(output_path)
        except Exception as e:
            import pandas as pd
            df = pd.DataFrame(rows, columns=columns)
            df.to_excel(output_path, index=False)

        return output_path

    def create_csv(self, columns: List[str], rows: List[List[Any]], filename: Optional[str] = None) -> str:
        """
        Generate a CSV dataset file.
        """
        fname = filename or f"generated_{uuid.uuid4().hex[:8]}.csv"
        if not fname.endswith(".csv"):
            fname += ".csv"
        output_path = os.path.join(settings.DOCUMENT_STORAGE_DIR, fname)

        import pandas as pd
        df = pd.DataFrame(rows, columns=columns)
        df.to_csv(output_path, index=False, encoding="utf-8-sig")
        return output_path

    def create_json(self, data: Any, filename: Optional[str] = None) -> str:
        """
        Generate a structured JSON file.
        """
        fname = filename or f"generated_{uuid.uuid4().hex[:8]}.json"
        if not fname.endswith(".json"):
            fname += ".json"
        output_path = os.path.join(settings.DOCUMENT_STORAGE_DIR, fname)

        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        return output_path

    async def ai_draft_and_create(
        self,
        prompt: str,
        file_type: str,
        title: Optional[str] = None,
        model: str = "gpt-4o"
    ) -> Dict[str, Any]:
        """
        Use LLM to draft content according to file type, then create the physical file.
        """
        file_type = file_type.lower().replace(".", "")
        doc_title = title or f"{file_type.upper()} Dokümanı"

        if file_type in ["pdf", "docx", "doc"]:
            # Prompt for structured document
            ai_prompt = f"""Kullanıcı şu konuda bir {file_type.upper()} dokümanı oluşturmak istiyor:
İSTEM: {prompt}
BAŞLIK: {doc_title}

Lütfen bu doküman için profesyonel, detaylı, eksiksiz bir içerik üret ve aşağıdaki JSON formatında döndür:
{{
  "title": "{doc_title}",
  "sections": [
    {{
      "title": "Bölüm Başlığı",
      "content": "Bölüm içeriği, açıklamalar, paragraflar...",
      "table": [
        ["Sütun 1", "Sütun 2", "Sütun 3"],
        ["Değer 1", "Değer 2", "Değer 3"]
      ]
    }}
  ]
}}
SADECE JSON döndür."""

            res = await llm_gateway.generate_response(
                messages=[
                    {"role": "system", "content": "You are a professional enterprise document authoring system. Return ONLY valid JSON."},
                    {"role": "user", "content": ai_prompt}
                ],
                model=model,
                temperature=0.4
            )

            raw_json = res.get("content", "").strip()
            if "```json" in raw_json:
                raw_json = raw_json.split("```json")[1].split("```")[0].strip()
            elif "```" in raw_json:
                raw_json = raw_json.split("```")[1].split("```")[0].strip()

            try:
                parsed = json.loads(raw_json)
                sections = parsed.get("sections", [])
                final_title = parsed.get("title", doc_title)
            except Exception:
                sections = [{"title": "Giriş ve Genel Bakış", "content": res.get("content", prompt)}]
                final_title = doc_title

            clean_filename = f"{final_title.lower().replace(' ', '_')[:30]}_{uuid.uuid4().hex[:6]}.{file_type}"
            if file_type == "pdf":
                file_path = self.create_pdf(final_title, sections, filename=clean_filename)
            else:
                file_path = self.create_docx(final_title, sections, filename=clean_filename)

            return {
                "filename": clean_filename,
                "file_type": file_type,
                "file_path": file_path,
                "title": final_title,
                "sections": sections
            }

        elif file_type in ["xlsx", "xls", "csv"]:
            ai_prompt = f"""Kullanıcı şu konuda bir {file_type.upper()} tablosu/veriseti oluşturmak istiyor:
İSTEM: {prompt}
BAŞLIK: {doc_title}

Lütfen bu tablo için gerçekçi, tutarlı, zengin bir veri seti üret ve aşağıdaki JSON formatında döndür:
{{
  "title": "{doc_title}",
  "columns": ["Sütun1", "Sütun2", "Sütun3", "Sütun4"],
  "rows": [
    ["Veri 1A", "Veri 1B", 1000, "Aktif"],
    ["Veri 2A", "Veri 2B", 2500, "Beklemede"]
  ]
}}
SADECE JSON döndür."""

            res = await llm_gateway.generate_response(
                messages=[
                    {"role": "system", "content": "You are a senior data engineer. Return ONLY valid JSON."},
                    {"role": "user", "content": ai_prompt}
                ],
                model=model,
                temperature=0.3
            )

            raw_json = res.get("content", "").strip()
            if "```json" in raw_json:
                raw_json = raw_json.split("```json")[1].split("```")[0].strip()
            elif "```" in raw_json:
                raw_json = raw_json.split("```")[1].split("```")[0].strip()

            try:
                parsed = json.loads(raw_json)
                columns = parsed.get("columns", ["Kolon1", "Kolon2", "Kolon3"])
                rows = parsed.get("rows", [["Örnek 1", "Örnek 2", 100]])
                final_title = parsed.get("title", doc_title)
            except Exception:
                columns = ["Kalem", "Miktar", "Birim Fiyat", "Toplam"]
                rows = [["Ürün A", 10, 150, 1500], ["Ürün B", 5, 200, 1000]]
                final_title = doc_title

            clean_filename = f"{final_title.lower().replace(' ', '_')[:30]}_{uuid.uuid4().hex[:6]}.{file_type}"
            if file_type in ["xlsx", "xls"]:
                file_path = self.create_excel(final_title, columns, rows, filename=clean_filename)
            else:
                file_path = self.create_csv(columns, rows, filename=clean_filename)

            return {
                "filename": clean_filename,
                "file_type": file_type,
                "file_path": file_path,
                "title": final_title,
                "columns": columns,
                "rows_count": len(rows)
            }

        else:  # JSON
            ai_prompt = f"""Kullanıcı şu konuda yapılandırılmış bir JSON dosyası oluşturmak istiyor:
İSTEM: {prompt}
BAŞLIK: {doc_title}

Lütfen profesyonel, hiyerarşik ve geçerli bir JSON yapısı üret. SADECE JSON formatında döndür."""

            res = await llm_gateway.generate_response(
                messages=[
                    {"role": "system", "content": "You are a software data architect. Return ONLY valid JSON."},
                    {"role": "user", "content": ai_prompt}
                ],
                model=model,
                temperature=0.3
            )

            raw_json = res.get("content", "").strip()
            if "```json" in raw_json:
                raw_json = raw_json.split("```json")[1].split("```")[0].strip()
            elif "```" in raw_json:
                raw_json = raw_json.split("```")[1].split("```")[0].strip()

            try:
                parsed = json.loads(raw_json)
            except Exception:
                parsed = {"title": doc_title, "content": prompt, "status": "generated"}

            clean_filename = f"{doc_title.lower().replace(' ', '_')[:30]}_{uuid.uuid4().hex[:6]}.json"
            file_path = self.create_json(parsed, filename=clean_filename)

            return {
                "filename": clean_filename,
                "file_type": "json",
                "file_path": file_path,
                "title": doc_title,
                "data": parsed
            }

document_creator_service = DocumentCreatorService()
